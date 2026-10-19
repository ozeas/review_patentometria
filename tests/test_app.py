"""Testes de integração da aplicação Django (fluxo completo com as amostras)."""
import io
import zipfile
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook

from studies import export, models as m, services, tabular

FIX = Path(__file__).parent / "fixtures"


class AppFlowTests(TestCase):
    def setUp(self):
        self.study = m.Study.objects.create(name="Teste")
        self.rnd = m.Round.objects.create(study=self.study, label="r1")
        for name in ("lens_sample.csv", "espacenet_sample.csv", "google_sample.csv", "derwent_sample.txt"):
            services.ingest(self.rnd, SimpleUploadedFile(name, (FIX / name).read_bytes()))
        # recorte BR: mesma amostra do Google com rótulo BR
        services.ingest(self.rnd, SimpleUploadedFile("google_BR.csv", (FIX / "google_sample.csv").read_bytes()))
        services.run_round(self.rnd)

    def test_ingest_and_run(self):
        self.assertEqual(self.rnd.files.filter(recorte="BR").count(), 1)
        self.assertEqual(self.rnd.families.count(), 5)
        p = services.compute(self.rnd).prisma()
        self.assertEqual(p["identified"], 15)
        self.assertEqual(p["pending_stage1"], 5)

    def test_decision_flow_and_reuse_across_rounds(self):
        fam = self.rnd.families.get(anchor="CN101413931")
        services.save_decision(self.study, fam, stage1="incluir", stage2="incluir",
                               stage2_reason="Exploração", stratum="C1")
        p = services.compute(self.rnd).prisma()
        self.assertEqual(p["final"], 1)
        rnd2 = m.Round.objects.create(study=self.study, label="r2")
        services.ingest(rnd2, SimpleUploadedFile("lens.csv", (FIX / "lens_sample.csv").read_bytes()))
        services.run_round(rnd2)
        self.assertEqual(services.compute(rnd2).prisma()["final"], 1)

    def test_dictionary_and_pages(self):
        services.upsert_applicant(self.study, "大庆油田有限责任公司", "DAQING OILFIELD CO LTD", "CNPC")
        services.upsert_applicant(self.study, "DAQING OILFIELD CO LTD", "DAQING OILFIELD CO LTD", "CNPC")
        services.run_round(self.rnd)
        fam = self.rnd.families.get(anchor="CN101413931")
        self.assertEqual(fam.fields["groups"], ["CNPC"])
        urls = [f"/", f"/s/{self.study.id}/", f"/s/{self.study.id}/config/", f"/r/{self.rnd.id}/",
                f"/r/{self.rnd.id}/familias/", f"/r/{self.rnd.id}/f/{fam.key}/", f"/r/{self.rnd.id}/ligacoes/",
                f"/r/{self.rnd.id}/depositantes/", f"/r/{self.rnd.id}/inventores/",
                f"/r/{self.rnd.id}/indicadores/", f"/r/{self.rnd.id}/indicadores/?conjunto=unificada",
                f"/r/{self.rnd.id}/vosviewer/", f"/s/{self.study.id}/tesauro/", f"/r/{self.rnd.id}/importar/",
                f"/r/{self.rnd.id}/familias/?status=pendente1&q=chromat"]
        for u in urls:
            r = self.client.get(u)
            self.assertEqual(r.status_code, 200, u)
        r = self.client.post(f"/r/{self.rnd.id}/f/{fam.key}/", {
            "action": "decision", "stage1": "incluir", "stage2": "excluir", "stage2_reason": "Downstream",
            "go": "next"})
        self.assertEqual(r.status_code, 302)
        d = services.decision_for_family(self.study, fam.numbers)
        self.assertEqual((d.stage1, d.stage2, d.stage2_reason), ("incluir", "excluir", "Downstream"))

    def test_exports(self):
        fam = self.rnd.families.get(anchor="CN101413931")
        services.save_decision(self.study, fam, stage1="incluir", stage2="incluir",
                               stage2_reason="Exploração", stratum="C1")
        m.ThesaurusEntry.objects.create(study=self.study, label="method", replace_by="")
        m.ThesaurusEntry.objects.create(study=self.study, label="gc-ms", replace_by="gc ms")
        wb = load_workbook(io.BytesIO(export.workbook_bytes(self.rnd)))
        for s in ("PRISMA", "Indicadores", "S2_Registros", "S2_Consolidacao", "S5_Base_Final", "S6_Estratos"):
            self.assertIn(s, wb.sheetnames)
        self.assertEqual(wb["S5_Base_Final"].max_row, 2)
        z = zipfile.ZipFile(io.BytesIO(export.vos_package(self.rnd)))
        self.assertEqual(sorted(z.namelist()), ["LEIA-ME.txt", "controle.csv", "corpus.txt", "scores.txt",
                                                "thesaurus.txt"])
        self.assertTrue(z.read("thesaurus.txt").decode().strip().endswith("gc ms"))
        self.assertEqual(len(z.read("corpus.txt").decode().strip().splitlines()), 1)

    def test_vos_import(self):
        mp = SimpleUploadedFile("map.txt", (
            "id\tlabel\tx\ty\tcluster\tweight<Links>\tweight<Total link strength>\tweight<Occurrences>\t"
            "score<Ano de prioridade>\tscore<Janela recente>\n"
            "1\tgas chromatography\t0\t0\t1\t2\t4\t10\t2015.5\t0.2\n"
            "2\tcrude oil\t0\t0\t1\t2\t4\t8\t2018.0\t0.3\n"
            "3\tdownhole\t0\t0\t2\t1\t1\t5\t2012.0\t0.1\n").encode())
        nw = SimpleUploadedFile("network.txt", b"1\t2\t3\n1\t3\t1\n")
        r = self.client.post(f"/r/{self.rnd.id}/vosviewer/", {"label": "teste", "map_file": mp,
                                                              "network_file": nw, "params": "VOSviewer 1.6.21"})
        self.assertEqual(r.status_code, 302)
        run = self.rnd.vos_runs.get()
        self.assertEqual(len(run.result["clusters"]), 2)
        self.assertEqual(run.result["clusters"][0]["coesao"], 0.75)
        r = self.client.post(f"/r/{self.rnd.id}/vosviewer/", {"action": "labels", "run": run.id,
                                                              "label_1": "Análise de fluidos"})
        run.refresh_from_db()
        self.assertEqual(run.cluster_labels["1"], "Análise de fluidos")
        self.assertEqual(self.client.get(f"/r/{self.rnd.id}/vosviewer/").status_code, 200)
        wb = load_workbook(io.BytesIO(export.workbook_bytes(self.rnd)))
        self.assertIn("S11_VOSviewer_clusters", wb.sheetnames)

    def test_import_decisions_csv(self):
        csv_text = "Números;Etapa 1;Motivo etapa 1;Etapa 2;Segmento/motivo etapa 2;Estrato\n" \
                   "WO 1999/058951 A1;incluir;;excluir;Downstream;\n" \
                   "US2016168990A1;excluir;Falso positivo lexical;;;\n"
        rows = tabular.read_table(SimpleUploadedFile("d.csv", csv_text.encode("utf-8")))
        r = tabular.import_decisions(self.study, rows)
        self.assertEqual(r["criadas"], 2)
        p = services.compute(self.rnd).prisma()
        self.assertEqual(p["stage1_excluded"], 1)
        self.assertEqual(p["stage2_excluded"], 1)
