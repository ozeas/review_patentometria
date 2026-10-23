"""Integração HTTP: registro dos fluxos, importação, edição de ajustes e exportações."""
import csv as _csv
import io
import tempfile
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from studies import export, models as m, services, tabular

FIX = Path(__file__).parent / "fixtures"


def _formula_prisma(wb, descricao):
    ws = wb["PRISMA"]
    for row in ws.iter_rows(min_col=2, max_col=3):
        if row[0].value == descricao:
            return row[1].value
    raise AssertionError(f"linha '{descricao}' não encontrada na aba PRISMA")


def _ingest_4(rnd):
    for name in ("lens_sample.csv", "espacenet_sample.csv", "google_sample.csv", "derwent_sample.txt"):
        services.ingest(rnd, SimpleUploadedFile(name, (FIX / name).read_bytes()))


class PrismaFormulaTests(TestCase):
    def setUp(self):
        self.study = m.Study.objects.create(name="Formula")
        self.rnd = m.Round.objects.create(study=self.study, label="r1")
        _ingest_4(self.rnd)
        services.run_round(self.rnd)

    def test_formulas_prisma_referenciam_linhas_corretas(self):
        wb = load_workbook(io.BytesIO(export.workbook_bytes(self.rnd)))
        self.assertEqual(_formula_prisma(wb, "Registros elegíveis à deduplicação"), "=C2-C8")
        self.assertEqual(_formula_prisma(wb, "Registros redundantes"), "=C9-C10")


class FluxoCriacaoTests(TestCase):
    def test_criar_rodada_via_http(self):
        study = m.Study.objects.create(name="Fluxo")
        r = self.client.post(f"/s/{study.id}/", {"label": "Rodada X",
                                                 "collection_date": "2026-10-15", "notes": ""})
        self.assertEqual(r.status_code, 302)
        rnd = study.rounds.get(label="Rodada X")
        self.assertEqual(rnd.collection_date.isoformat(), "2026-10-15")


class UploadHttpTests(TestCase):
    def test_upload_multiplo_query_e_concepts(self):
        study = m.Study.objects.create(name="Up")
        rnd = m.Round.objects.create(study=study, label="r1")
        files = [
            SimpleUploadedFile("lens_nucleo.csv", (FIX / "lens_sample.csv").read_bytes()),
            SimpleUploadedFile("espacenet_nucleo.csv", (FIX / "espacenet_sample.csv").read_bytes()),
            SimpleUploadedFile("lens_nucleo_query.txt", b"G01N30 AND oil"),
            SimpleUploadedFile("google_concepts.csv", b"conceito ignorado"),
        ]
        r = self.client.post(f"/r/{rnd.id}/upload/", {"files": files})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(rnd.files.count(), 2)          # concepts ignorado
        sf = rnd.files.get(source="lens")
        self.assertEqual(sf.query_text, "G01N30 AND oil")
        self.assertGreater(rnd.records.count(), 0)

    def test_executar_http_grava_familias(self):
        study = m.Study.objects.create(name="Run")
        rnd = m.Round.objects.create(study=study, label="r1")
        _ingest_4(rnd)
        r = self.client.post(f"/r/{rnd.id}/executar/")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(rnd.families.count(), 5)


class ArquivoCrudTests(TestCase):
    def test_editar_e_remover_arquivo(self):
        study = m.Study.objects.create(name="Arq")
        rnd = m.Round.objects.create(study=study, label="r1")
        sf = services.ingest(rnd, SimpleUploadedFile("lens_nucleo.csv",
                                                      (FIX / "lens_sample.csv").read_bytes()))
        r = self.client.post(f"/arquivo/{sf.id}/", {"query_text": "X", "executed_at": "hoje"})
        self.assertEqual(r.status_code, 302)
        sf.refresh_from_db()
        self.assertEqual(sf.query_text, "X")
        r = self.client.post(f"/arquivo/{sf.id}/remover/")
        self.assertEqual(r.status_code, 302)
        self.assertFalse(m.SourceFile.objects.filter(pk=sf.id).exists())
        self.assertEqual(m.Record.objects.filter(source_file=sf).count(), 0)


class ImportarRodadaTests(TestCase):
    def test_importa_pasta_com_query_e_ignora_concepts(self):
        d = Path(tempfile.mkdtemp())
        (d / "lens_nucleo.csv").write_bytes((FIX / "lens_sample.csv").read_bytes())
        (d / "espacenet_nucleo.csv").write_bytes((FIX / "espacenet_sample.csv").read_bytes())
        (d / "lens_nucleo_query.txt").write_text("G01N30 AND oil", encoding="utf-8")
        (d / "google_concepts.csv").write_text("ignorar", encoding="utf-8")
        (d / ".DS_Store").write_text("ignorar", encoding="utf-8")
        out = io.StringIO()
        call_command("importar_rodada", estudo="Cmd", rodada="r1", pasta=str(d), stdout=out)
        study = m.Study.objects.get(name="Cmd")
        rnd = study.rounds.get(label="r1")
        self.assertEqual(rnd.files.count(), 2)                       # concepts e .DS_Store ignorados
        self.assertEqual(rnd.files.get(source="lens").query_text, "G01N30 AND oil")
        self.assertIn("lens_nucleo.csv", out.getvalue())


class LigacoesCrudTests(TestCase):
    def setUp(self):
        self.study = m.Study.objects.create(name="Lig")
        self.rnd = m.Round.objects.create(study=self.study, label="r1")
        _ingest_4(self.rnd)
        services.run_round(self.rnd)

    def test_aceitar_e_remover_ligacao(self):
        r = self.client.post(f"/r/{self.rnd.id}/ligacoes/",
                             {"action": "accept", "number_a": "US1", "number_b": "US2"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.study.manual_links.count(), 1)
        ml = self.study.manual_links.get()
        r = self.client.post(f"/r/{self.rnd.id}/ligacoes/", {"action": "delete", "id": ml.id})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.study.manual_links.count(), 0)


class DicionariosCrudTests(TestCase):
    def test_depositantes_via_http(self):
        study = m.Study.objects.create(name="Dep")
        rnd = m.Round.objects.create(study=study, label="r1")
        r = self.client.post(f"/r/{rnd.id}/depositantes/",
                             {"raw_1": "FOO CORP", "h_1": "FOO CORP", "g_1": "FOO"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(study.applicant_aliases.count(), 1)
        a = study.applicant_aliases.get()
        self.assertEqual((a.raw_example, a.harmonized, a.group), ("FOO CORP", "FOO CORP", "FOO"))

    def test_inventores_via_http(self):
        study = m.Study.objects.create(name="Inv")
        rnd = m.Round.objects.create(study=study, label="r1")
        r = self.client.post(f"/r/{rnd.id}/inventores/",
                             {"raw_1": "张三", "s_1": "ZHANG SAN", "ok_1": "on"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(study.inventor_aliases.count(), 1)


class TesauroCrudTests(TestCase):
    def setUp(self):
        self.study = m.Study.objects.create(name="Tes")

    def test_adicionar_e_excluir_entrada(self):
        r = self.client.post(f"/s/{self.study.id}/tesauro/",
                             {"label": "gas chrom", "replace_by": "", "category": "", "note": ""})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.study.thesaurus.count(), 1)
        e = self.study.thesaurus.get()
        r = self.client.post(f"/s/{self.study.id}/tesauro/", {"action": "delete", "id": e.id})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.study.thesaurus.count(), 0)

    def test_importar_tesauro(self):
        conteudo = "label\treplace by\ngas chromatograph\t\nGC\tgas chromatography\n"
        f = SimpleUploadedFile("thes.txt", conteudo.encode("utf-8"))
        r = self.client.post(f"/s/{self.study.id}/tesauro/", {"action": "import", "file": f})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.study.thesaurus.count(), 2)


class ImportDataTests(TestCase):
    def test_importar_decisoes_via_http(self):
        study = m.Study.objects.create(name="Imp")
        rnd = m.Round.objects.create(study=study, label="r1")
        csv_text = "Números;Etapa 1\nWO 1999/058951 A1;incluir\n"
        f = SimpleUploadedFile("d.csv", csv_text.encode("utf-8"))
        r = self.client.post(f"/r/{rnd.id}/importar/", {"kind": "decisoes", "file": f})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(study.decisions.count(), 1)
        self.assertEqual(study.decisions.get().stage1, "incluir")


class TriagemTests(TestCase):
    def test_triagem_redireciona_primeira_pendente(self):
        study = m.Study.objects.create(name="Tri")
        rnd = m.Round.objects.create(study=study, label="r1")
        _ingest_4(rnd)
        services.run_round(rnd)
        r = self.client.get(f"/r/{rnd.id}/triagem/1/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/f/", r["Location"])


class ExportHttpTests(TestCase):
    def setUp(self):
        self.study = m.Study.objects.create(name="Exp")
        self.rnd = m.Round.objects.create(study=self.study, label="r1")
        _ingest_4(self.rnd)
        services.run_round(self.rnd)
        fam = self.rnd.families.first()
        services.save_decision(self.study, fam, stage1="incluir", stage2="incluir",
                               stage2_reason="Segmento incluído", stratum="S1")

    def test_downloads_respondem(self):
        casos = [
            (f"/r/{self.rnd.id}/planilha.xlsx",
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            (f"/r/{self.rnd.id}/decisoes.csv", "text/csv"),
            (f"/r/{self.rnd.id}/vosviewer/pacote.zip", "application/zip"),
            (f"/s/{self.study.id}/tesauro/thesaurus.txt", "text/plain"),
            (f"/s/{self.study.id}/exportar/depositantes.csv", "text/csv"),
            (f"/s/{self.study.id}/exportar/inventores.csv", "text/csv"),
            (f"/s/{self.study.id}/exportar/ligacoes.csv", "text/csv"),
        ]
        for url, ct in casos:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
            self.assertIn(ct, r["Content-Type"], url)
            self.assertGreater(len(r.content), 0, url)

    def test_decisoes_csv_roundtrip(self):
        r = self.client.get(f"/r/{self.rnd.id}/decisoes.csv")
        rows = list(_csv.DictReader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        self.assertGreaterEqual(len(rows), 1)
        antes = self.study.decisions.count()
        tabular.import_decisions(self.study, rows)
        self.assertEqual(self.study.decisions.count(), antes)

    def test_exportar_planilha_command(self):
        d = Path(tempfile.mkdtemp()) / "base.xlsx"
        call_command("exportar_planilha", self.rnd.id, saida=str(d), stdout=io.StringIO())
        self.assertTrue(d.exists())
        self.assertGreater(d.stat().st_size, 0)
        wb = load_workbook(d)
        self.assertIn("PRISMA", wb.sheetnames)
