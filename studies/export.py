"""Geração da planilha base (fonte única dos números do estudo).

As abas de registros, famílias e decisões trazem os dados; as abas PRISMA e
Estratos usam fórmulas sobre essas abas, de modo que qualquer contagem do
texto possa ser rastreada até a linha que a origina.
"""
from __future__ import annotations

import io
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.prevalence import vos_text
from core.records import SOURCE_LABELS, SOURCES
from core.vosviewer import build_corpus, build_thesaurus

from . import models as m
from .services import compute

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
J = "; "


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, start=1):
        w = (widths or {}).get(h) or min(60, max(10, len(str(h)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _cell(v):
    if isinstance(v, (list, tuple, set)):
        return J.join(str(x) for x in v)
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v


def _dec(f):
    d = f["decision"]
    if not d:
        return ["", "", "", "", "", ""]
    return [d.stage1, d.stage1_reason, d.stage2, d.stage2_reason, d.stratum, d.notes]


def build_workbook(round_: m.Round) -> Workbook:
    result = compute(round_)
    cfg = result.config
    cons = result.consolidation
    study = round_.study
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------ LEIA-ME
    ws = wb.create_sheet("LEIA-ME")
    lines = [
        ("Estudo", study.name), ("Rodada", round_.label),
        ("Data da coleta", cfg.get("collection_date")),
        ("Gerado em", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Unidade de contagem", "família consolidada (encadeamento de números de publicação/prioridade)"),
        ("Regra de prevalência", "; ".join(f"{k}: {' > '.join(v)}" for k, v in cfg["prevalence"].items())),
        ("Prioridades no encadeamento", f"somente números com ≥ {cfg['prio_link_min_digits']} dígitos"),
        ("Modelos de utilidade excluídos", "registros cujas publicações têm todas código de tipo "
                                           + "/".join(cfg["utility_kinds"])),
        ("Períodos", "; ".join(f"{a}–{b}" for a, b in cfg["periods"])),
        ("Janela recente", f"{cfg['window_months']} meses antes da coleta (data da 1ª publicação)"),
        ("", ""),
        ("Abas", ""),
        ("S1_Estrategia", "arquivos de exportação, consultas registradas e contagens por base"),
        ("S2_Registros", "todos os registros identificados, normalizados, com a família atribuída"),
        ("S2_Consolidacao", "famílias consolidadas com origem de cada campo e decisões"),
        ("S2_Ligacoes_manuais", "ligações manuais aplicadas e não aplicadas"),
        ("S2_Grupos_familia / S2_Inventores / S2_IPC", "formato longo (família × valor)"),
        ("S3_Etapa1 / S4_Etapa2 / S5_Base_Final", "elegibilidade em duas etapas e base final"),
        ("S6_Estratos", "contagens e indicadores por estrato (fórmulas)"),
        ("S7_Situacao_juridica / S8_Citacoes / S9_Rede_citacoes", "análises de anterioridade"),
        ("S10_VOS_corpus", "controle do corpus do VOSviewer (linha ↔ família)"),
        ("S11_VOSviewer_clusters", "indicadores por cluster (quando há mapa importado)"),
        ("S12_Recorte", "recortes comparativos (ex.: BR) e correspondência com o núcleo"),
        ("PRISMA", "fluxo de identificação, triagem, elegibilidade e inclusão (fórmulas)"),
        ("Indicadores", "indicadores descritivos da base final"),
    ]
    for a, b in lines:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 110
    for c in ws["A"]:
        c.font = Font(bold=True)

    # ------------------------------------------------------------ S1
    rows = []
    for sf in round_.files.all():
        rows.append([SOURCE_LABELS.get(sf.source, sf.source), sf.recorte, sf.original_name, sf.n_records,
                     sf.executed_at, sf.query_text])
    _sheet(wb, "S1_Estrategia", ["Base", "Recorte", "Arquivo", "Registros", "Execução", "Consulta"],
           rows, {"Consulta": 100, "Arquivo": 30})

    # ------------------------------------------------------------ S2_Registros
    excluded = {r.record_id for r in cons.excluded_utility}
    rows = []
    for r in cons.identified:
        rows.append([r.record_id, r.source, r.recorte, r.row,
                     [p.raw for p in r.pubs], r.pub_numbers, [p.kind for p in r.pubs],
                     "sim" if r.record_id in excluded else "não",
                     cons.record_family.get(r.record_id, ""),
                     r.title or r.dwpi_title, r.earliest_priority.isoformat() if r.earliest_priority else ""])
    _sheet(wb, "S2_Registros", ["Registro", "Base", "Recorte", "Linha", "Números exportados",
                                "Números normalizados", "Códigos de tipo", "Modelo de utilidade",
                                "Família", "Título", "Prioridade mais antiga"], rows,
           {"Números exportados": 40, "Números normalizados": 40, "Título": 60})

    # ------------------------------------------------------------ S2_Consolidacao
    rows = []
    for f in result.families:
        x, o = f["fields"], f["origin"]
        rows.append([
            f["key"], x["n_records"], x["n_sources"], *[("sim" if s in x["sources"] else "não") for s in SOURCES],
            x["publications"], o.get("publications"), x["title"], o.get("title"), x["abstract_kind"],
            x["ipc"], o.get("ipc"), x["applicants"], x["groups"], o.get("applicants"), x["inventors"],
            o.get("inventors"), x["earliest_priority"], x["priority_year"], x["period"],
            x["origin_jurisdiction"], o.get("origin_jurisdiction"), x["origin_priority_number"],
            x["first_publication"], x["recent_window"], x["legal_status"], x["cited_by_max"],
            x["lexical_level"], f["family"].manual_link_ids, *_dec(f),
        ])
    _sheet(wb, "S2_Consolidacao", [
        "Família", "Registros", "Bases", *[SOURCE_LABELS[s] for s in SOURCES],
        "Publicações", "Origem publicações", "Título", "Origem título", "Resumo",
        "IPC", "Origem IPC", "Depositantes", "Grupos econômicos", "Origem depositantes", "Inventores",
        "Origem inventores", "Prioridade mais antiga", "Ano de prioridade", "Período",
        "Jurisdição de origem", "Fonte da jurisdição", "Prioridade usada", "1ª publicação",
        "Janela recente", "Situação jurídica", "Citações recebidas (máx.)", "Nível lexical",
        "Ligações manuais", "Etapa 1", "Motivo etapa 1", "Etapa 2", "Segmento/motivo etapa 2",
        "Estrato", "Observações"], rows, {"Título": 60, "Publicações": 40, "Depositantes": 40,
                                         "Inventores": 40, "IPC": 30})

    # ------------------------------------------------------------ ligações
    applied = {i for f in result.families for i in f["family"].manual_link_ids}
    rows = [[ml.id, ml.number_a, ml.number_b, ml.reason, "sim" if ml.id in applied else "não"]
            for ml in study.manual_links.all()]
    _sheet(wb, "S2_Ligacoes_manuais", ["ID", "Número A", "Número B", "Justificativa", "Aplicada nesta rodada"],
           rows, {"Justificativa": 80})

    _sheet(wb, "S2_Grupos_familia", ["Família", "Grupo econômico", "Depositantes harmonizados", "Grafias exportadas"],
           [[f["key"], g, f["fields"]["applicants"], f["fields"]["applicants_raw"]]
            for f in result.families for g in f["fields"]["groups"]], {"Grupo econômico": 40})
    _sheet(wb, "S2_Inventores", ["Família", "Inventor (padronizado)"],
           [[f["key"], i] for f in result.families for i in f["fields"]["inventors"]], {"Inventor (padronizado)": 40})
    _sheet(wb, "S2_IPC", ["Família", "IPC", "Subclasse", "Grupo principal"],
           [[f["key"], c, c[:4], c.split("/")[0]] for f in result.families for c in f["fields"]["ipc"]])

    # ------------------------------------------------------------ elegibilidade
    def _brief(f):
        x = f["fields"]
        return [f["key"], x["title"], x["groups"], x["priority_year"], x["origin_jurisdiction"]]

    _sheet(wb, "S3_Etapa1", ["Família", "Título", "Grupos", "Ano", "Jurisdição", "Etapa 1", "Motivo"],
           [[*_brief(f), *(_dec(f)[:2])] for f in result.families], {"Título": 70})
    st1 = result.stage("etapa1")
    _sheet(wb, "S4_Etapa2", ["Família", "Título", "Grupos", "Ano", "Jurisdição", "Etapa 2", "Segmento/motivo"],
           [[*_brief(f), *(_dec(f)[2:4])] for f in st1], {"Título": 70})
    final = result.stage("final")
    _sheet(wb, "S5_Base_Final", ["Família", "Título", "Grupos", "Ano", "Jurisdição", "Período", "Segmento",
                                 "Estrato", "Janela recente", "Situação jurídica", "Nível lexical"],
           [[*_brief(f), f["fields"]["period"], f["decision"].stage2_reason, f["decision"].stratum,
             f["fields"]["recent_window"], f["fields"]["legal_status"], f["fields"]["lexical_level"]]
            for f in final], {"Título": 70})

    # ------------------------------------------------------------ S6 estratos (fórmulas)
    ws = wb.create_sheet("S6_Estratos")
    ws.append(["Estrato", "Definição", "Famílias", "% da base final"])
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    n_final_ref = "COUNTA(S5_Base_Final!A:A)-1"
    for s in cfg["strata"]:
        r = ws.max_row + 1
        ws.append([s["code"], f"{s['name']}: {s['definition']}",
                   f'=COUNTIF(S5_Base_Final!H:H,"{s["code"]}")', f"=IFERROR(100*C{r}/({n_final_ref}),0)"])
    ws.append([])
    ws.append(["Período × estrato"] + [s["code"] for s in cfg["strata"]])
    for a, b in cfg["periods"]:
        lab = f"{a}–{b}"
        ws.append([lab] + [f'=COUNTIFS(S5_Base_Final!F:F,"{lab}",S5_Base_Final!H:H,"{s["code"]}")'
                           for s in cfg["strata"]])
    ws.append([])
    ws.append(["Nível lexical × estrato"] + [s["code"] for s in cfg["strata"]])
    for lev in ("titulo", "titulo+resumo", "ausente"):
        ws.append([lev] + [f'=COUNTIFS(S5_Base_Final!K:K,"{lev}",S5_Base_Final!H:H,"{s["code"]}")'
                           for s in cfg["strata"]])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    # ------------------------------------------------------------ anterioridade
    _sheet(wb, "S7_Situacao_juridica", ["Família", "Situação (agregada)", "Situações Lens"],
           [[f["key"], f["fields"]["legal_status"], f["fields"]["legal_status_raw"]]
            for f in final if f["fields"]["has_lens"]])
    _sheet(wb, "S8_Citacoes", ["Família", "Título", "Citações recebidas (máx. entre publicações)"],
           [[f["key"], f["fields"]["title"], f["fields"]["cited_by_max"]]
            for f in sorted(final, key=lambda f: -(f["fields"]["cited_by_max"] or -1)) if f["fields"]["has_lens"]],
           {"Título": 70})
    from core.indicators import internal_citation_edges
    edges = internal_citation_edges([{"key": f["key"], "fields": f["fields"]} for f in final])
    _sheet(wb, "S9_Rede_citacoes", ["Família citante", "Família citada"], edges)

    # ------------------------------------------------------------ VOS
    corpus, scores, control = build_corpus(final, vos_text)
    _sheet(wb, "S10_VOS_corpus", ["Linha do corpus", "Família", "Texto usado", "Ano", "Janela"],
           [[c["linha"], c["familia"], c["modo"], c["ano"], c["janela"]] for c in control])
    vr = round_.vos_runs.first()
    if vr and vr.result:
        keys = vr.result.get("score_keys", [])
        rows = [[c["cluster"], vr.cluster_labels.get(str(c["cluster"]), ""), c["termos"], c["ocorrencias"],
                 *[c.get(k) for k in keys], c["coesao"], c["termos_lista"]] for c in vr.result["clusters"]]
        _sheet(wb, "S11_VOSviewer_clusters", ["Cluster", "Rótulo", "Termos", "Ocorrências", *keys, "Coesão",
                                              "Termos (por ocorrência)"], rows, {"Termos (por ocorrência)": 90})

    # ------------------------------------------------------------ recorte
    if result.br:
        _sheet(wb, "S12_Recorte", ["Família do recorte", "Bases", "Números", "Famílias do núcleo",
                                   "Na base final", "Depositantes"],
               [[r["familia_recorte"], r["fontes"], r["numeros"], r["familias_nucleo"], r["na_base_final"],
                 r["depositantes"]] for r in result.br["families"]], {"Números": 50, "Depositantes": 50})

    # ------------------------------------------------------------ PRISMA (fórmulas)
    p = result.prisma()
    ws = wb.create_sheet("PRISMA", 1)
    ws.append(["Etapa", "Descrição", "n", "Conferência (valor calculado pela aplicação)"])
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    core_rec = cfg["core_recorte"]
    ws.append(["Identificação", "Registros identificados (núcleo)",
               f'=COUNTIF(S2_Registros!C:C,"{core_rec}")', p["identified"]])
    for s in SOURCES:
        ws.append(["Identificação", f"  {SOURCE_LABELS[s]}",
                   f'=COUNTIFS(S2_Registros!B:B,"{s}",S2_Registros!C:C,"{core_rec}")',
                   p["identified_by_source"][s]])
    ws.append(["Triagem", "Modelos de utilidade excluídos",
               f'=COUNTIFS(S2_Registros!H:H,"sim",S2_Registros!C:C,"{core_rec}")', p["excluded_utility"]])
    ws.append(["Triagem", "Registros elegíveis à deduplicação", "=C2-C7", p["eligible"]])
    ws.append(["Triagem", "Base unificada (famílias)", "=COUNTA(S2_Consolidacao!A:A)-1", p["families"]])
    ws.append(["Triagem", "Registros redundantes", "=C8-C9", p["redundant"]])
    ws.append(["Elegibilidade", "Excluídas na etapa 1", '=COUNTIF(S3_Etapa1!F:F,"excluir")', p["stage1_excluded"]])
    for reason in cfg["stage1"]["exclusion_reasons"]:
        ws.append(["Elegibilidade", f"  {reason}", f'=COUNTIFS(S3_Etapa1!F:F,"excluir",S3_Etapa1!G:G,"{reason}")',
                   p["stage1_excluded_by_reason"].get(reason, 0)])
    ws.append(["Elegibilidade", "Mantidas na etapa 1", '=COUNTIF(S3_Etapa1!F:F,"incluir")', p["stage1_included"]])
    ws.append(["Elegibilidade", "Excluídas na etapa 2", '=COUNTIF(S4_Etapa2!F:F,"excluir")', p["stage2_excluded"]])
    for reason in cfg["stage2"]["exclusion_reasons"]:
        ws.append(["Elegibilidade", f"  {reason}", f'=COUNTIFS(S4_Etapa2!F:F,"excluir",S4_Etapa2!G:G,"{reason}")',
                   p["stage2_excluded_by_reason"].get(reason, 0)])
    ws.append(["Inclusão", "Base final", "=COUNTA(S5_Base_Final!A:A)-1", p["final"]])
    for seg in cfg["stage2"]["inclusion_segments"]:
        ws.append(["Inclusão", f"  {seg}", f'=COUNTIF(S5_Base_Final!G:G,"{seg}")', p["final_by_segment"].get(seg, 0)])
    for s in cfg["strata"]:
        ws.append(["Inclusão", f"  Estrato {s['code']} ({s['name']})", f'=COUNTIF(S5_Base_Final!H:H,"{s["code"]}")',
                   p["strata"].get(s["code"], 0)])
    ws.append(["Pendências", "Famílias sem decisão na etapa 1", "", p["pending_stage1"]])
    ws.append(["Pendências", "Famílias da etapa 1 sem decisão na etapa 2", "", p["pending_stage2"]])
    for col, w in zip("ABCD", (16, 52, 10, 40)):
        ws.column_dimensions[col].width = w

    # ------------------------------------------------------------ Indicadores
    ind = result.indicators("final")
    ws = wb.create_sheet("Indicadores", 2)
    bold = Font(bold=True)

    def block(title, note, header, rows):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        if note:
            ws.append([note])
        ws.append(header)
        for c in ws[ws.max_row]:
            c.font = bold
        for r in rows:
            ws.append(r)
        ws.append([])

    n = ind["n"]
    block("Períodos de prioridade", f"Denominador: {ind['periodos']['denominador']} famílias com prioridade (de {n}).",
          ["Período", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in ind["periodos"]["linhas"]])
    block("Jurisdição de origem", f"Denominador: {ind['jurisdicao_origem']['denominador']} famílias.",
          ["Jurisdição", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in ind["jurisdicao_origem"]["linhas"]])
    c = ind["grupos"]["concentracao"]
    block("Grupos econômicos (atribuição a todos os grupos da família)",
          f"Denominador: {n} famílias; percentuais não somam 100%. CR4 {c['CR4']}%; CR8 {c['CR8']}%; "
          f"HHI {c['HHI']}; Gini {c['Gini']} ({c['grupos']} grupos; {c['atribuicoes']} atribuições).",
          ["Grupo", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in ind["grupos"]["todas"]])
    block("Subclasses IPC", f"Denominador: {ind['ipc_subclasses']['denominador']} famílias com IPC.",
          ["Subclasse", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in ind["ipc_subclasses"]["linhas"]])
    block("Grupos principais IPC", f"Denominador: {ind['ipc_grupos']['denominador']} famílias com IPC.",
          ["Grupo", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in ind["ipc_grupos"]["linhas"]])
    if "bloco_classificacao" in ind:
        block("Bloco de classificação nas IPC exportadas",
              f"Denominador: {n} famílias; {ind['bloco_classificacao']['sem_ipc']} sem IPC exportada.",
              ["Código", "Famílias com o código", "Famílias só com este código do bloco"],
              [[r["codigo"], r["n_com_codigo"], r["n_exclusivo"]] for r in ind["bloco_classificacao"]["linhas"]])
    jr = ind["janela_recente"]
    block("Janela recente", f"1ª publicação nos {cfg['window_months']} meses anteriores à coleta.",
          ["Famílias", "Denominador", "%"], [[jr["n"], jr["denominador"], jr["pct"]]])
    sj = ind["situacao_juridica"]
    block("Situação jurídica (Lens)", f"Denominador: {sj['denominador']} famílias com registro na Lens; "
                                      "situação mais favorável entre as publicações (limite superior para vigentes).",
          ["Situação", "Famílias", "%"], [[r["categoria"], r["n"], r["pct"]] for r in sj["linhas"]])
    ci = ind["citacoes"]
    block("Citações recebidas (Lens, máximo por família)", "", ["Medida", "Valor"],
          [[k, v] for k, v in ci.items()])
    block("Rede de citações internas (Derwent)", "", ["Arestas", "Famílias envolvidas"],
          [[ind["rede_citacoes_internas"]["arestas"], ind["rede_citacoes_internas"]["familias_envolvidas"]]])
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    return wb


def workbook_bytes(round_: m.Round) -> bytes:
    bio = io.BytesIO()
    build_workbook(round_).save(bio)
    return bio.getvalue()


def vos_package(round_: m.Round) -> bytes:
    result = compute(round_)
    final = result.stage("final")
    corpus, scores, control = build_corpus(final, vos_text)
    thes = build_thesaurus([{"label": t.label, "replace_by": t.replace_by} for t in round_.study.thesaurus.all()])
    readme = (
        "Pacote de entrada do VOSviewer\n"
        f"Estudo: {round_.study.name} | Rodada: {round_.label}\n"
        f"Famílias da base final: {len(final)}\n\n"
        "corpus.txt  — um documento por linha (título + resumo original; sem resumo original,\n"
        "              título + seções NOVELTY e USE do DWPI; senão, apenas o título)\n"
        "scores.txt  — escores por documento (mesma ordem do corpus)\n"
        "thesaurus.txt — tesauro (unificações e exclusões; exclusões antes das unificações)\n"
        "controle.csv  — linha do corpus ↔ família\n\n"
        "No VOSviewer: Create map → Text data → Read data from VOSviewer corpus file;\n"
        "informar corpus.txt e scores.txt; contagem binária; aplicar thesaurus.txt.\n"
    )
    ctrl = "linha;familia;modo;ano;janela\n" + "\n".join(
        f"{c['linha']};{c['familia']};{c['modo']};{c['ano']};{c['janela']}" for c in control) + "\n"
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("corpus.txt", corpus)
        z.writestr("scores.txt", scores)
        z.writestr("thesaurus.txt", thes)
        z.writestr("controle.csv", ctrl)
        z.writestr("LEIA-ME.txt", readme)
    return bio.getvalue()
