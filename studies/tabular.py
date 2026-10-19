"""Importação e exportação tabular (CSV/XLSX) de decisões, dicionários e
ligações manuais — permite reaproveitar o trabalho manual de uma planilha
anterior e de outras rodadas."""
from __future__ import annotations

import csv
import io
import re
import unicodedata

from core.normalize import normalize_pub
from core.vosviewer import parse_thesaurus

from . import models as m
from .services import upsert_applicant, upsert_inventor

ALIASES = {
    "numbers": ["numeros", "numero", "numeros normalizados", "publicacoes", "anchor", "ancora",
                "publication number", "numero de publicacao"],
    "stage1": ["etapa 1", "etapa1", "stage1", "decisao etapa 1"],
    "stage1_reason": ["motivo etapa 1", "motivo 1", "motivo1", "stage1_reason", "motivo"],
    "stage2": ["etapa 2", "etapa2", "stage2", "decisao etapa 2"],
    "stage2_reason": ["segmento/motivo etapa 2", "segmento", "motivo etapa 2", "motivo 2", "motivo2",
                      "stage2_reason", "segmento/motivo"],
    "stratum": ["estrato", "stratum", "categoria"],
    "notes": ["observacoes", "obs", "notas", "notes"],
    "raw": ["grafia exportada", "grafia", "raw", "nome exportado", "original"],
    "harmonized": ["nome harmonizado", "harmonizado", "harmonized", "nome padronizado"],
    "group": ["grupo economico", "grupo", "group"],
    "standard": ["forma padronizada", "padronizado", "standard", "pinyin"],
    "number_a": ["numero a", "number_a", "a"],
    "number_b": ["numero b", "number_b", "b"],
    "reason": ["justificativa", "reason"],
}


def _norm(h: str) -> str:
    h = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode().lower()
    h = re.sub(r"\(.*?\)", "", h)
    return re.sub(r"\s+", " ", h).strip()


def read_table(upload) -> list[dict]:
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "") for h in rows[0]]
        return [dict(zip(header, ["" if v is None else str(v) for v in r])) for r in rows[1:]]
    text = data.decode("utf-8-sig", errors="replace")
    delim = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delim))


def pick(row: dict, field: str) -> str:
    keys = {_norm(k): k for k in row}
    for alias in [field] + ALIASES.get(field, []):
        k = keys.get(_norm(alias))
        if k is not None:
            return (row[k] or "").strip()
    return ""


def _numbers(s: str) -> list[str]:
    out = []
    for chunk in re.split(r"[;,|\n]+", s or ""):
        chunk = chunk.strip()
        if not chunk:
            continue
        p = normalize_pub(chunk)
        if p and p.country:
            out.append(p.number)
            continue
        for tok in chunk.split():
            p = normalize_pub(tok)
            if p and p.country:
                out.append(p.number)
    return sorted(set(out))


def _yes_no(v: str) -> str:
    v = _norm(v)
    if v in ("incluir", "incluida", "mantida", "manter", "sim", "include", "in", "1", "true"):
        return "incluir"
    if v in ("excluir", "excluida", "nao", "exclude", "out", "0", "false"):
        return "excluir"
    return ""


def import_decisions(study: m.Study, rows: list[dict]) -> dict:
    created = updated = skipped = 0
    existing = list(study.decisions.all())
    for row in rows:
        nums = _numbers(pick(row, "numbers"))
        if not nums:
            skipped += 1
            continue
        d = next((x for x in existing if set(x.numbers) & set(nums)), None)
        if d is None:
            d = m.Decision(study=study, anchor=nums[0], numbers=nums)
            existing.append(d)
            created += 1
        else:
            d.numbers = sorted(set(d.numbers) | set(nums))
            updated += 1
        d.stage1 = _yes_no(pick(row, "stage1")) or d.stage1
        d.stage1_reason = pick(row, "stage1_reason") or d.stage1_reason
        d.stage2 = _yes_no(pick(row, "stage2")) or d.stage2
        d.stage2_reason = pick(row, "stage2_reason") or d.stage2_reason
        d.stratum = pick(row, "stratum") or d.stratum
        d.notes = pick(row, "notes") or d.notes
        d.save()
    return {"criadas": created, "atualizadas": updated, "ignoradas": skipped}


def export_decisions_csv(study: m.Study, families=None) -> str:
    """Modelo de decisões: uma linha por família (quando ``families`` é dado)
    ou por decisão registrada."""
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Família", "Números", "Título", "Etapa 1", "Motivo etapa 1", "Etapa 2",
                "Segmento/motivo etapa 2", "Estrato", "Observações"])
    from .services import decisions_map
    if families is not None:
        dm = decisions_map(study, families)
        for f in families:
            d = dm.get(f.key)
            w.writerow([f.key, " ".join(f.numbers), f.title, *(
                [d.stage1, d.stage1_reason, d.stage2, d.stage2_reason, d.stratum, d.notes] if d else [""] * 6)])
    else:
        for d in study.decisions.all():
            w.writerow(["", " ".join(d.numbers), "", d.stage1, d.stage1_reason, d.stage2, d.stage2_reason,
                        d.stratum, d.notes])
    return out.getvalue()


def import_applicants(study, rows) -> int:
    n = 0
    for row in rows:
        raw, h = pick(row, "raw"), pick(row, "harmonized")
        if raw and h:
            upsert_applicant(study, raw, h, pick(row, "group"))
            n += 1
    return n


def import_inventors(study, rows) -> int:
    n = 0
    for row in rows:
        raw, s = pick(row, "raw"), pick(row, "standard")
        if raw and s:
            upsert_inventor(study, raw, s)
            n += 1
    return n


def import_links(study, rows) -> int:
    n = 0
    for row in rows:
        a, b = _numbers(pick(row, "number_a")), _numbers(pick(row, "number_b"))
        if a and b:
            m.ManualLink.objects.get_or_create(study=study, number_a=a[0], number_b=b[0],
                                               defaults={"reason": pick(row, "reason") or ""})
            n += 1
    return n


def import_thesaurus(study, text: str, replace: bool = False) -> int:
    entries = parse_thesaurus(text)
    if replace:
        study.thesaurus.all().delete()
    for e in entries:
        m.ThesaurusEntry.objects.update_or_create(study=study, label=e["label"],
                                                  defaults={"replace_by": e["replace_by"]})
    return len(entries)


def export_table_csv(header, rows) -> str:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(header)
    w.writerows(rows)
    return out.getvalue()
